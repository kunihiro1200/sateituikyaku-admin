import openpyxl, warnings
warnings.filterwarnings("ignore")

wb = openpyxl.load_workbook("C:/Users/kunih/Downloads/beppu_tuugaku2.xlsx")
ws = wb["R8"]

print("Max columns:", ws.max_column)
print("Max rows:", ws.max_row)
print()

# ヘッダー行（Row 0-2）を全列表示
for i in range(3):
    row = [ws.cell(i+1, c).value for c in range(1, ws.max_column+1)]
    print("Row %d:" % i, [str(v)[:20] if v else None for v in row])

print()
print("データ行サンプル（5件）:")
for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True)):
    if i >= 5:
        break
    vals = [str(v)[:25] if v else None for v in row[:ws.max_column]]
    # Noneを除いて表示
    non_none = [(j, v) for j, v in enumerate(vals) if v]
    print("  Row %d:" % (i+3), non_none)
