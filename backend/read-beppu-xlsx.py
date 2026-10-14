import openpyxl

for fname in ["beppu_tuugaku2.xlsx", "beppu_06020150.xlsx"]:
    path = "C:/Users/kunih/Downloads/" + fname
    try:
        wb = openpyxl.load_workbook(path)
        print("\n=== %s ===" % fname)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print("Sheet: %s (%d rows)" % (sheet_name, ws.max_row))
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 6:
                    break
                print("  Row %d:" % i, [str(c)[:30] if c is not None else None for c in row[:6]])
    except Exception as e:
        print("Error %s: %s" % (fname, e))
